import io
from datetime import date

from docx import Document
from fastapi.testclient import TestClient

from src.api.main import app

client = TestClient(app)


def test_health():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_version():
    response = client.get("/version")
    assert response.status_code == 200
    body = response.json()
    assert "version" in body
    assert "environment" in body


def test_extract_endpoint():
    response = client.post("/extract", json={"text": "ООО Ромашка против Иванова"})
    assert response.status_code == 200
    data = response.json()
    # На HTTP-границе ExtractionResult сериализуется обратно в JSON/dict.
    assert "plaintiff" in data
    assert "case_number" in data


def test_extract_result_feeds_generate_claim(sample_ruling_text):
    """Сквозной сценарий продукта: разобрали решение — сгенерировали иск.

    Раньше ответ /extract нельзя было передать в /generate/claim как есть:
    сумма отдавалась только строкой в том виде, в каком стояла в тексте
    («500 000»), а /generate/claim принимает claim_amount как float, то есть
    клиенту приходилось нормализовывать её самому. Тест проверяет именно
    стык двух эндпоинтов, а не разбор суммы отдельно.
    """
    extracted = client.post("/extract", json={"text": sample_ruling_text}).json()
    assert extracted["claim_amount_value"] == 500000.0

    generated = client.post(
        "/generate/claim",
        json={
            "plaintiff": extracted["plaintiff"],
            "defendant": extracted["defendant"],
            "claim_amount": extracted["claim_amount_value"],
            "case_number": extracted["case_number"],
        },
    )
    assert generated.status_code == 200
    assert generated.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def test_extract_endpoint_empty_text_returns_422():
    # min_length=1 на уровне Pydantic-схемы отклоняет пустую строку раньше,
    # чем запрос вообще доходит до EntityExtractor.
    response = client.post("/extract", json={"text": ""})
    assert response.status_code == 422


def test_extract_endpoint_missing_field_returns_422():
    response = client.post("/extract", json={})
    assert response.status_code == 422


def test_generate_claim_endpoint():
    response = client.post(
        "/generate/claim",
        json={
            "plaintiff": "ООО Тест",
            "defendant": "Иванов",
            "claim_amount": 100000,
            "case_number": "А40-1",
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in response.headers["content-disposition"]


def test_generate_appeal_endpoint():
    response = client.post(
        "/generate/appeal",
        json={
            "plaintiff": "ООО Тест",
            "defendant": "Иванов",
            "claim_amount": 100000,
            "case_number": "А40-1",
        },
    )
    assert response.status_code == 200


def test_generate_contract_endpoint():
    response = client.post(
        "/generate/contract",
        json={
            "customer": "ООО Тест",
            "contractor": "Иванов",
            "contract_amount": 100000,
            "contract_number": "Д-1",
        },
    )
    assert response.status_code == 200


def test_generate_contract_uses_todays_date_not_a_hardcoded_one():
    # Регрессия: раньше contract_date был захардкожен как "01.07.2026"
    # независимо от реальной даты генерации.
    response = client.post(
        "/generate/contract",
        json={
            "customer": "ООО Тест",
            "contractor": "Иванов",
            "contract_amount": 100000,
            "contract_number": "Д-1",
        },
    )
    assert response.status_code == 200

    doc = Document(io.BytesIO(response.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    today_str = date.today().strftime("%d.%m.%Y")
    assert today_str in full_text
    assert "01.07.2026" not in full_text


def test_generate_contract_accepts_custom_contractor_representative():
    response = client.post(
        "/generate/contract",
        json={
            "customer": "ООО Тест",
            "contractor": "Иванов",
            "contract_amount": 100000,
            "contract_number": "Д-1",
            "contractor_representative": "Сидоров С.С., по доверенности",
        },
    )
    assert response.status_code == 200

    doc = Document(io.BytesIO(response.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Сидоров С.С., по доверенности" in full_text
    assert "Иванов И.И." not in full_text


def test_generate_contract_rejects_missing_fields():
    response = client.post("/generate/contract", json={"customer": "ООО Тест"})
    assert response.status_code == 422


def test_generate_claim_rejects_non_positive_amount():
    response = client.post(
        "/generate/claim",
        json={
            "plaintiff": "ООО Тест",
            "defendant": "Иванов",
            "claim_amount": 0,
            "case_number": "А40-1",
        },
    )
    assert response.status_code == 422


def test_generate_claim_rejects_missing_fields():
    response = client.post("/generate/claim", json={"plaintiff": "ООО Тест"})
    assert response.status_code == 422


def test_404_for_unknown_route():
    response = client.get("/does-not-exist")
    assert response.status_code == 404


def test_error_responses_never_leak_stack_trace():
    # Отправляем валидный JSON, но с типом, который не пройдёт валидацию —
    # ответ должен быть структурированным JSON, а не HTML/traceback.
    response = client.post("/extract", json={"text": 12345})
    assert response.status_code == 422
    assert response.headers["content-type"].startswith("application/json")


def test_extract_rejected_without_api_key_when_auth_enabled(monkeypatch):
    from src.api import security

    monkeypatch.setattr(security.settings, "api_keys", "expected-key")
    response = client.post("/extract", json={"text": "ООО Ромашка против Иванова"})
    assert response.status_code == 401


def test_extract_allowed_with_valid_api_key(monkeypatch):
    from src.api import security

    monkeypatch.setattr(security.settings, "api_keys", "expected-key")
    response = client.post(
        "/extract",
        json={"text": "ООО Ромашка против Иванова"},
        headers={"X-API-Key": "expected-key"},
    )
    assert response.status_code == 200


def test_unhandled_exception_never_leaks_internal_details():
    from src.api.dependencies import get_generator

    secret_detail = "internal server error detail"

    class _BoomGenerator:
        def generate(self, template_name, data):
            raise RuntimeError(secret_detail)

        app.dependency_overrides[get_generator] = lambda: _BoomGenerator()
    local_client = TestClient(app, raise_server_exceptions=False)
    try:
        response = local_client.post(
            "/generate/claim",
            json={
                "plaintiff": "ООО Тест",
                "defendant": "Иванов",
                "claim_amount": 100000,
                "case_number": "А40-1",
            },
        )
    finally:
        app.dependency_overrides.pop(get_generator, None)


    assert response.status_code == 500
    assert secret_detail not in response.text
    assert response.json() == {"error": "Internal server error"}
